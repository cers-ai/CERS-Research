"""报告导出：Markdown 解析为简单的 DOCX / PDF 文档。"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

CJK_FONT_PATH = "C:/Windows/FONTS/simhei.ttf"


def export_report(payload: dict[str, Any]) -> dict[str, Any]:
    content = str(payload.get("content") or "")
    output_format = str(payload.get("format") or "").lower()
    output_path = Path(str(payload.get("path") or ""))
    title = str(payload.get("title") or "研究报告").strip() or "研究报告"

    if not content.strip():
        return {"ok": False, "message": "报告正文为空，无法导出。"}
    if output_format not in ("docx", "pdf"):
        return {"ok": False, "message": "不支持的报告导出格式。"}
    if not output_path.parent:
        return {"ok": False, "message": "导出路径无效。"}

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        if output_format == "docx":
            _write_docx(content, title, output_path)
        else:
            _write_pdf(content, title, output_path)
    except Exception as exc:  # noqa: BLE001 - 导出失败需要返回可读原因
        return {"ok": False, "message": f"报告导出失败：{type(exc).__name__}"}

    return {"ok": True, "path": str(output_path), "format": output_format}


def _parse_blocks(markdown: str) -> list[tuple[str, Any]]:
    blocks: list[tuple[str, Any]] = []
    lines = markdown.replace("\r\n", "\n").split("\n")
    index = 0
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            blocks.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
            index += 1
            continue

        if stripped.startswith("|") and stripped.count("|") >= 2:
            table_rows: list[list[str]] = []
            while index < len(lines):
                row = lines[index].strip()
                if not row.startswith("|"):
                    break
                cells = [cell.strip() for cell in row.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells if cell):
                    index += 1
                    continue
                table_rows.append([_plain(cell) for cell in cells])
                index += 1
            if table_rows:
                blocks.append(("table", table_rows))
            continue

        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        if bullet:
            blocks.append(("bullet", _plain(bullet.group(1))))
            index += 1
            continue

        blocks.append(("text", _plain(stripped)))
        index += 1

    return blocks


def _plain(value: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r"\1（\2）", value)
    text = re.sub(r"\*\*|__|\*|_|`", "", text)
    return re.sub(r"\s+", " ", text).strip()


def _write_docx(markdown: str, title: str, path: Path) -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    document.add_heading(title, level=0)
    for kind, value in _parse_blocks(markdown):
        if kind == "heading":
            level, text = value
            document.add_heading(text, level=min(int(level), 4))
        elif kind == "bullet":
            document.add_paragraph(str(value), style="List Bullet")
        elif kind == "table":
            rows: list[list[str]] = value
            columns = max(len(row) for row in rows)
            table = document.add_table(rows=0, cols=columns)
            table.style = "Table Grid"
            for row in rows:
                cells = table.add_row().cells
                for index in range(columns):
                    cells[index].text = row[index] if index < len(row) else ""
        else:
            document.add_paragraph(str(value))

    document.save(str(path))


def _write_pdf(markdown: str, title: str, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font_name = "Helvetica"
    font_path = Path(CJK_FONT_PATH)
    if font_path.is_file():
        try:
            pdfmetrics.registerFont(TTFont("CJK", str(font_path)))
            font_name = "CJK"
        except Exception:  # noqa: BLE001 - 字体不可用时退回默认字体
            font_name = "Helvetica"

    base = ParagraphStyle(
        "Base",
        fontName=font_name,
        fontSize=10.5,
        leading=16,
        spaceAfter=6,
        wordWrap="CJK",
    )
    heading = ParagraphStyle(
        "Heading",
        parent=base,
        fontSize=15,
        leading=22,
        spaceBefore=10,
        spaceAfter=8,
    )
    small_heading = ParagraphStyle(
        "SmallHeading",
        parent=base,
        fontSize=12,
        leading=18,
        spaceBefore=8,
        spaceAfter=6,
    )
    title_style = ParagraphStyle(
        "Title",
        parent=base,
        fontSize=20,
        leading=28,
        spaceAfter=14,
    )
    cell_style = ParagraphStyle(
        "Cell",
        parent=base,
        fontSize=8.5,
        leading=12,
        spaceAfter=0,
    )

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=title,
    )

    story: list[Any] = [Paragraph(_escape(title), title_style), Spacer(1, 8)]
    for kind, value in _parse_blocks(markdown):
        if kind == "heading":
            level, text = value
            style = heading if int(level) <= 2 else small_heading
            story.append(Paragraph(_escape(text), style))
        elif kind == "bullet":
            story.append(Paragraph(f"• {_escape(str(value))}", base))
        elif kind == "table":
            rows: list[list[str]] = value
            columns = max(len(row) for row in rows)
            data = [
                [
                    Paragraph(_escape(row[index] if index < len(row) else ""), cell_style)
                    for index in range(columns)
                ]
                for row in rows
            ]
            available = A4[0] - 44 * mm
            table = Table(
                data,
                colWidths=[available / columns] * columns,
                repeatRows=1,
            )
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CDD6E3")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph(_escape(str(value)), base))

    document.build(story)


def _escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
